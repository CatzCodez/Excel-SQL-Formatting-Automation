import os
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

RAW_DIR = "raw_reports"
OUT_DIR = "cleaned_reports"
os.makedirs(OUT_DIR, exist_ok=True)

def snake_case(text):
    if pd.isna(text):
        return ""
    text = str(text).strip()
    text = text.lower()
    text = re.sub(r"[^\w\s]", "", text)
    text = re.sub(r"\s+", "_", text)
    return text

def detect_header_row(df):
    for i in range(min(10, len(df))):
        if df.iloc[i].astype(str).str.contains("year", case=False).any():
            return i
    return 0

def process_file(src_path, dst_path):
    ext = os.path.splitext(src_path)[1].lower()
    df0 = pd.read_csv(src_path, header=None, dtype=str) if ext == ".csv" else pd.read_excel(src_path, header=None, dtype=str)

    header_row = detect_header_row(df0)
    df = df0.iloc[header_row + 1:].copy()
    df.columns = df0.iloc[header_row]
    df.columns = [snake_case(col) for col in df.columns]

    df = df.loc[:, ~df.columns.str.startswith("unnamed")]
    df.dropna(how="all", inplace=True)

    first_col = df.columns[0]
    df[first_col] = df[first_col].apply(snake_case)

    for col in df.columns[1:]:
        df[col] = df[col].astype(str).str.replace(",", "").str.strip()
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0).astype(int)

    df.to_excel(dst_path, index=False, sheet_name="Data")

    wb = load_workbook(dst_path)
    ws = wb["Data"]
    table = Table(displayName="DataTable", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    ws.add_table(table)
    wb.save(dst_path)
    print(f"✅ Cleaned: {os.path.basename(src_path)}")

# Run it
for file in os.listdir(RAW_DIR):
    if file.lower().endswith((".csv", ".xlsx")):
        src = os.path.join(RAW_DIR, file)
        dst = os.path.join(OUT_DIR, os.path.splitext(file)[0] + "_formatted.xlsx")
        if os.path.exists(dst):
            os.remove(dst)
        process_file(src, dst)
